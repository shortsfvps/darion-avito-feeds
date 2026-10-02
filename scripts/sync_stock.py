"""
Модуль загрузки и проверки актуальных остатков и цен из Google Sheets.
Выполняет:
1. Загрузку файла остатков по экспортной ссылке Google Таблицы через requests.
2. Проверку HTTP-статуса и обработку прав доступа Google Drive (публичный vs приватный).
3. Чтение листа «Остатки СПб» через pandas (из буфера/кэша).
4. Анализ названий и индексов колонок (включая проверку столбцов J и K: цена и остаток).
5. Сверку артикулов между номенклатурой (data/Продукты (19).xlsx) и остатками.
"""

import os
import sys
import io
import requests
import pandas as pd
import openpyxl

# Обеспечиваем корректный вывод UTF-8 в Windows консоль
if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

GOOGLE_SHEET_URL = "https://docs.google.com/spreadsheets/d/1ooUhgWojjoce0YGbzrmzJtL9tz9xm_OU0QBSaqEMCas/export?format=xlsx"
DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
LOCAL_CACHE_PATH = os.path.join(DATA_DIR, "remains_spb_cache.xlsx")
NOMENCLATURE_PATH = os.path.join(DATA_DIR, "Продукты (19).xlsx")


def get_nomenclature_skus():
    """Извлекает все непустые артикулы из номенклатурного файла."""
    if not os.path.exists(NOMENCLATURE_PATH):
        return []
    wb = openpyxl.load_workbook(NOMENCLATURE_PATH, data_only=True)
    skus = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for r in range(3, ws.max_row + 1):
            sku = ws.cell(r, 1).value
            if sku is not None and str(sku).strip():
                skus.append(str(sku).strip())
    return skus


def fetch_google_sheets(url=GOOGLE_SHEET_URL, cache_path=LOCAL_CACHE_PATH):
    """
    Загружает файл остатков через requests.
    Возвращает (content_bytes, status_code, message)
    """
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    try:
        response = requests.get(url, headers=headers, timeout=20)
        status = response.status_code

        # Проверяем, получен ли реальный Excel-файл (сигнатура PK..)
        is_excel = response.content[:2] == b"PK" and len(response.content) > 1000

        if status == 200 and is_excel:
            with open(cache_path, "wb") as f:
                f.write(response.content)
            return response.content, 200, "Файл успешно загружен из Google Sheets"
        elif status == 401 or status == 403 or not is_excel:
            msg = (
                f"Google Sheets вернул HTTP {status} (Требуется авторизация / Ограниченный доступ). "
                "В настройках доступа таблицы необходимо включить: 'Все, у кого есть ссылка -> Читатель'."
            )
            return None, status, msg
        else:
            return None, status, f"Неожиданный ответ сервера: HTTP {status}"
    except Exception as e:
        return None, -1, f"Ошибка сети при запросе к Google Sheets: {str(e)}"


def check_stock_data(excel_bytes=None, file_path=None):
    """
    Считывает лист «Остатки СПб» через pandas и проверяет колонки и сопоставление артикулов.
    """
    source = io.BytesIO(excel_bytes) if excel_bytes else file_path

    if not source or (isinstance(source, str) and not os.path.exists(source)):
        return {
            "success": False,
            "error": "Отсутствует источник данных для чтения"
        }

    try:
        # Проверяем список листов
        xl = pd.ExcelFile(source)
        sheet_names = xl.sheet_names
        target_sheet = "Остатки СПб"

        if target_sheet not in sheet_names:
            # Если точного совпадения нет, ищем похожий лист
            matched = [s for s in sheet_names if "спб" in s.lower() or "остат" in s.lower()]
            if matched:
                target_sheet = matched[0]
            else:
                target_sheet = sheet_names[0]

        df = pd.read_excel(source, sheet_name=target_sheet)

        # Анализ колонок
        columns = list(df.columns)
        num_cols = len(columns)

        # Столбец J (10-й по счету, индекс 9)
        # Столбец K (11-й по счету, индекс 10)
        col_j_name = columns[9] if num_cols > 9 else None
        col_k_name = columns[10] if num_cols > 10 else None

        # Ищем столбец артикулов
        sku_col = None
        for col in columns:
            col_l = str(col).lower()
            if "артикул" in col_l or "код" in col_l or "арт" in col_l:
                sku_col = col
                break
        if sku_col is None and num_cols > 0:
            sku_col = columns[0]

        # Извлекаем артикулы из остатков
        stock_skus = set()
        if sku_col is not None:
            stock_skus = set(df[sku_col].dropna().astype(str).str.strip())

        # Сверяем с номенклатурой
        nom_skus = get_nomenclature_skus()
        matched_skus = [sku for sku in nom_skus if sku in stock_skus]

        return {
            "success": True,
            "sheet_used": target_sheet,
            "all_sheets": sheet_names,
            "total_rows": len(df),
            "total_cols": num_cols,
            "columns": columns,
            "col_j_price": {
                "letter": "J",
                "index_0_based": 9,
                "name": col_j_name
            },
            "col_k_stock": {
                "letter": "K",
                "index_0_based": 10,
                "name": col_k_name
            },
            "sku_column": sku_col,
            "nom_skus_count": len(nom_skus),
            "matched_skus_count": len(matched_skus),
            "matched_skus": matched_skus,
            "sample_rows": df.head(3).to_dict(orient="records")
        }
    except Exception as e:
        return {
            "success": False,
            "error": f"Ошибка обработки Excel-файла: {str(e)}"
        }


def run_stock_check():
    print("=" * 80)
    print("ПРОВЕРКА ЧТЕНИЯ ОСТАТКОВ ИЗ GOOGLE SHEETS")
    print("=" * 80)
    print(f"URL: {GOOGLE_SHEET_URL}")

    content, status, msg = fetch_google_sheets()
    print(f"Статус HTTP запроса: {status}")
    print(f"Результат: {msg}\n")

    if status == 200 and content:
        res = check_stock_data(excel_bytes=content)
    elif os.path.exists(LOCAL_CACHE_PATH) and os.path.getsize(LOCAL_CACHE_PATH) > 1000:
        print(f"Используем ранее сохраненный локальный кэш: {LOCAL_CACHE_PATH}")
        res = check_stock_data(file_path=LOCAL_CACHE_PATH)
    else:
        # Создаем демонстрационный / тестовый шаблон остатков на основе спецификации
        # Столбец J - Цена, Столбец K - Остаток, лист "Остатки СПб"
        print("Создаем эталонный файл структуры остатков data/remains_spb_template.xlsx...")
        test_path = os.path.join(DATA_DIR, "remains_spb_template.xlsx")
        
        # Создаем 11 колонок (A..K)
        # Col A (0): Артикул
        # Col B (1): Наименование
        # Col C (2): Ед. изм.
        # Col D (3): Производитель
        # Col E (4): Модель
        # Col F (5): Категория
        # Col G (6): Склад
        # Col H (7): Зарезервировано
        # Col I (8): Мин. заказ
        # Col J (9): Цена, руб.
        # Col K (10): Остаток СПб, шт.
        sample_data = {
            "Артикул": ["V9007", "VRN-001", "VRN-002"],
            "Наименование": [
                "Светодиодный светильник Virona 48Вт Универсальный",
                "Светильник Virona 36Вт",
                "Трековый светильник LED 20Вт"
            ],
            "Ед. изм.": ["шт", "шт", "шт"],
            "Производитель": ["VIRONA", "VIRONA", "VIRONA"],
            "Модель": ["VRN-UNE-48", "VRN-UNE-36", "VRN-TRK-20"],
            "Категория": ["LED Светильники", "LED Светильники", "Трековые системы"],
            "Склад": ["СПб Основной", "СПб Основной", "СПб Основной"],
            "Резерв": [0, 0, 0],
            "Мин. заказ": [1, 1, 1],
            "Цена, руб": [3450.00, 2800.00, 1950.00],        # Столбец J (индекс 9)
            "Остаток СПб, шт": [142, 28, 55]                # Столбец K (индекс 10)
        }
        sample_df = pd.DataFrame(sample_data)
        with pd.ExcelWriter(test_path, engine="openpyxl") as writer:
            sample_df.to_excel(writer, sheet_name="Остатки СПб", index=False)
        print(f"Эталонный шаблон структуры создан: {test_path}")
        res = check_stock_data(file_path=test_path)

    if res.get("success"):
        print(f"Лист: [{res['sheet_used']}] (всего листов в книге: {len(res['all_sheets'])})")
        print(f"Количество строк с данными: {res['total_rows']}")
        print(f"Количество колонок: {res['total_cols']}")
        print(f"Колонка артикула: [{res['sku_column']}]")
        print(f"Столбец J (индекс {res['col_j_price']['index_0_based']}): {res['col_j_price']['name']}")
        print(f"Столбец K (индекс {res['col_k_stock']['index_0_based']}): {res['col_k_stock']['name']}")
        print(f"Артикулов в номенклатуре: {res['nom_skus_count']}")
        print(f"Сопоставлено артикулов: {res['matched_skus_count']} ({res['matched_skus']})")
    else:
        print(f"Ошибка проверки остатков: {res.get('error')}")

    print("=" * 80)
    return res


if __name__ == "__main__":
    run_stock_check()
